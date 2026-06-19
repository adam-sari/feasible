"""Tabular fact ingestion (xlsx) — orders parsed into Jobs.

Tabular facts are *referenced* by typed rules, never re-expressed as logic.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from feasible.ir.domain import Job

ORDER_COLUMNS = ["id", "product", "family", "qty", "due_date", "priority"]


def load_orders_xlsx(path: str | Path, sheet: str = "orders") -> list[Job]:
    """Read an orders sheet (header row = ORDER_COLUMNS) into Jobs."""
    wb = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col = {name: header.index(name) for name in ORDER_COLUMNS if name in header}
    missing = [c for c in ORDER_COLUMNS if c not in col]
    if missing:
        raise ValueError(f"orders sheet missing columns: {missing}")
    jobs: list[Job] = []
    for row in rows[1:]:
        if row[col["id"]] is None:
            continue
        jobs.append(
            Job(
                id=str(row[col["id"]]),
                product=str(row[col["product"]]),
                family=str(row[col["family"]]),
                qty=int(row[col["qty"]]),
                due_date=int(row[col["due_date"]]),
                priority=int(row[col["priority"]]),
            )
        )
    return jobs


def write_orders_xlsx(path: str | Path, jobs: list[Job], sheet: str = "orders") -> None:
    """Render Jobs as an orders xlsx (used to build the realism-layer fixture)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append([c for c in ORDER_COLUMNS])
    for j in jobs:
        ws.append([j.id, j.product, j.family, j.qty, j.due_date, j.priority])
    wb.save(str(path))
